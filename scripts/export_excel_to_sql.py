#!/usr/bin/env python3
import glob
import os
import shutil
import unicodedata
import re
import openpyxl

PL_TO_ASCII = str.maketrans({
    'ą': 'a', 'ć': 'c', 'ę': 'e', 'ł': 'l', 'ń': 'n',
    'ó': 'o', 'ś': 's', 'ź': 'z', 'ż': 'z',
    'Ą': 'a', 'Ć': 'c', 'Ę': 'e', 'Ł': 'l', 'Ń': 'n',
    'Ó': 'o', 'Ś': 's', 'Ź': 'z', 'Ż': 'z',
})

DEFAULT_TDP_AILMENTS = [
    "hashimoto",
    "insulinooporność",
    "cukrzyca",
    "nietolerancja histaminy",
    "nadciśnienie tętnicze",
    "wysoki poziom cholesterolu",
    "wysoki poziom trójglicerydów",
    "lipoedema",
    "nadwaga/otyłość",
    "SIBO",
    "IMO",
    "niedoczynność tarczycy"
]


def generate_schema_sql(output_path, product_tables):
    """
    Dynamicznie generuje 01_init_schema.sql uwzględniając wszystkie wykryte tabele produktów.
    """
    lines = [
        "-- =========================================================",
        "-- Inicjalizacja schematu bazy danych Diet-Med (diet-med-DB)",
        "-- =========================================================",
        "",
        "-- 1. Tabela z listą dolegliwości (TDP)",
        "CREATE TABLE IF NOT EXISTS dolegliwosci (",
        "    id SERIAL PRIMARY KEY,",
        "    kod VARCHAR(100) NOT NULL UNIQUE",
        ");",
        ""
    ]

    for idx, (tbl, title) in enumerate(product_tables, start=2):
        lines.extend([
            f"-- {idx}. Tabela produktów dla {title}",
            f"DROP TABLE IF EXISTS {tbl} CASCADE;",
            f"CREATE TABLE IF NOT EXISTS {tbl} (",
            f"    id SERIAL PRIMARY KEY,",
            f"    rodzaj VARCHAR(255) NOT NULL UNIQUE,",
            f"    status VARCHAR(20) NOT NULL CHECK (status IN ('dozwolone', 'umiarkowane', 'zakazane')),",
            f"    ilosc NUMERIC(10, 2) NULL,",
            f"    jednostka VARCHAR(50) NULL,",
            f"    komentarz TEXT NULL",
            f");",
            f"CREATE UNIQUE INDEX IF NOT EXISTS idx_{tbl}_rodzaj ON {tbl}(rodzaj);",
            f"CREATE INDEX IF NOT EXISTS idx_{tbl}_status ON {tbl}(status);",
            ""
        ])

    lines.extend([
        f"-- Tabela zgłoszeń z formularza pacjentów",
        "CREATE TABLE IF NOT EXISTS zgloszenia (",
        "    id SERIAL PRIMARY KEY,",
        "    email VARCHAR(255) NULL,",
        "    dolegliwosci_ids INTEGER[] NOT NULL,",
        "    pdf_path VARCHAR(255) NULL,",
        "    status VARCHAR(50) DEFAULT 'wygenerowano',",
        "    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP",
        ");"
    ])

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Generated dynamic schema at {output_path} with {len(product_tables)} product tables")


def generate_tdp_sql(output_path, extra_ailments=None):
    """
    Generuje 02_seed_tdp.sql zawierający bazowe oraz dynamicznie dodane dolegliwości.
    """
    ailments = list(DEFAULT_TDP_AILMENTS)
    if extra_ailments:
        for a in extra_ailments:
            norm_existing = [x.lower() for x in ailments]
            if a.lower() not in norm_existing:
                ailments.append(a)

    lines = [
        "-- =========================================================",
        "-- Wypełnienie tabeli dolegliwości (TDP)",
        "-- =========================================================",
        "INSERT INTO dolegliwosci (kod) VALUES"
    ]
    vals = [f"('{kod}')" for kod in ailments]
    lines.append(",\n".join(vals) + " ON CONFLICT (kod) DO NOTHING;")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Generated {output_path} with {len(ailments)} ailments")


def generate_table_sql(excel_path, table_name, title_name, output_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    STATUS_PRIORITY = {"dozwolone": 1, "umiarkowane": 2, "zakazane": 3}
    products = {}

    for r in range(3, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        name = str(vals[0]).strip() if vals[0] is not None else ""
        if not name:
            continue

        # Jeśli tylko kolumna 1 ma wartość, to nagłówek kategorii
        if not any(vals[1:]):
            continue

        status = None
        if vals[1] and str(vals[1]).strip().lower() == "x":
            status = "dozwolone"
        elif vals[2] and str(vals[2]).strip().lower() == "x":
            status = "umiarkowane"
        elif vals[3] and str(vals[3]).strip().lower() == "x":
            status = "zakazane"
        elif any(vals[4:7]):
            # Jeśli produkt ma ilość/jednostkę/komentarz bez 'x', to z definicji warunkowość -> umiarkowane
            status = "umiarkowane"

        if not status:
            continue

        qty = vals[4]
        try:
            qty_val = float(qty) if qty is not None else None
        except (ValueError, TypeError):
            qty_val = None

        unit = str(vals[5]).strip() if vals[5] else None
        comment = str(vals[6]).strip() if vals[6] else None

        key = name.lower()
        if key not in products:
            products[key] = {
                "rodzaj": name,
                "status": status,
                "ilosc": qty_val,
                "jednostka": unit,
                "komentarze": [comment] if comment else []
            }
        else:
            existing = products[key]
            existing_score = STATUS_PRIORITY.get(existing["status"], 0)
            new_score = STATUS_PRIORITY.get(status, 0)

            # Wyższy priorytet (zakazane > umiarkowane > dozwolone)
            if new_score > existing_score:
                existing["status"] = status
                existing["ilosc"] = qty_val
                existing["jednostka"] = unit
                if comment and comment not in existing["komentarze"]:
                    existing["komentarze"].append(comment)
            elif new_score == existing_score and status == "umiarkowane":
                if qty_val is not None and existing["ilosc"] is not None:
                    if qty_val < existing["ilosc"]:
                        existing["ilosc"] = qty_val
                        existing["jednostka"] = unit
                elif qty_val is not None and existing["ilosc"] is None:
                    existing["ilosc"] = qty_val
                    existing["jednostka"] = unit
                if comment and comment not in existing["komentarze"]:
                    existing["komentarze"].append(comment)
            elif new_score == existing_score and status == "dozwolone":
                if comment and comment not in existing["komentarze"]:
                    existing["komentarze"].append(comment)

    rows = []
    for item in products.values():
        name_escaped = item["rodzaj"].replace("'", "''")
        qty_str = str(item["ilosc"]) if item["ilosc"] is not None else "NULL"

        if item["jednostka"]:
            u_esc = item["jednostka"].replace("'", "''")
            unit_str = f"'{u_esc}'"
        else:
            unit_str = "NULL"

        if item["komentarze"]:
            c_combined = " / ".join(item["komentarze"]).replace("'", "''")
            comm_str = f"'{c_combined}'"
        else:
            comm_str = "NULL"

        rows.append(f"('{name_escaped}', '{item['status']}', {qty_str}, {unit_str}, {comm_str})")

    lines = [
        f"-- =========================================================",
        f"-- Wypełnienie tabeli produktów {title_name} ({len(rows)} unikalnych wierszy)",
        f"-- =========================================================",
        f"DROP TABLE IF EXISTS {table_name} CASCADE;",
        f"CREATE TABLE IF NOT EXISTS {table_name} (",
        f"    id SERIAL PRIMARY KEY,",
        f"    rodzaj VARCHAR(255) NOT NULL UNIQUE,",
        f"    status VARCHAR(20) NOT NULL CHECK (status IN ('dozwolone', 'umiarkowane', 'zakazane')),",
        f"    ilosc NUMERIC(10, 2) NULL,",
        f"    jednostka VARCHAR(50) NULL,",
        f"    komentarz TEXT NULL",
        f");",
        f"CREATE UNIQUE INDEX IF NOT EXISTS idx_{table_name}_rodzaj ON {table_name}(rodzaj);",
        f"CREATE INDEX IF NOT EXISTS idx_{table_name}_status ON {table_name}(status);",
        f"",
        f"INSERT INTO {table_name} (rodzaj, status, ilosc, jednostka, komentarz) VALUES"
    ]
    lines.append(",\n".join(rows) + ";")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"Generated {output_path} with {len(rows)} unique products (0 duplicates)")
    return len(rows)


def parse_ailment_from_filename(path):
    fn = os.path.basename(path)
    fn_norm = unicodedata.normalize('NFC', fn)

    if 'baza' in fn_norm.lower() or fn_norm.lower() == 'tabela produktów.xlsx':
        return None

    m = re.search(r'tabela\s+produkt[oó]w\s+(.+)\.xlsx', fn_norm, re.IGNORECASE)
    if not m:
        return None

    raw = m.group(1).strip()
    raw = re.sub(r'-\d+$', '', raw).strip()

    if 'insulino' in raw.lower():
        code = 'insulinooporność'
        table_name = 'insulinoopornosc_produkty'
        title = 'Insulinooporność'
    elif 'sibo' in raw.lower():
        code = 'SIBO'
        table_name = 'sibo_produkty'
        title = 'SIBO'
    elif 'hashimoto' in raw.lower():
        code = 'hashimoto'
        table_name = 'hashimoto_produkty'
        title = 'Hashimoto'
    else:
        clean_ascii = raw.lower().translate(PL_TO_ASCII)
        clean_slug = re.sub(r'[^a-z0-9]+', '_', clean_ascii).strip('_')
        table_name = f'{clean_slug}_produkty'
        code = raw
        title = raw.capitalize()

    return {
        'code': code,
        'table_name': table_name,
        'title': title,
        'file_path': path,
        'is_dash1': '-1' in fn_norm
    }


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    root_dir = os.path.dirname(base_dir)

    all_xlsx = glob.glob(os.path.join(root_dir, "*.xlsx"))
    excel_files = [f for f in all_xlsx if not f.endswith('.bak') and not os.path.basename(f).startswith('~$')]

    discovered = {}
    for f in excel_files:
        info = parse_ailment_from_filename(f)
        if not info:
            continue
        code = info['code']
        # Prefer file without '-1' if both exist
        if code not in discovered or (discovered[code]['is_dash1'] and not info['is_dash1']):
            discovered[code] = info

    init_dir = os.path.join(base_dir, "init-scripts")
    backend_seeds_dir = os.path.join(root_dir, "diet-med-backend", "app", "seeds")
    os.makedirs(init_dir, exist_ok=True)

    # Wyczyść stare pliki seedów przed ponownym generowaniem
    for old_f in glob.glob(os.path.join(init_dir, "[0-9][0-9]_seed_*.sql")):
        os.remove(old_f)
    if os.path.exists(backend_seeds_dir):
        for old_f in glob.glob(os.path.join(backend_seeds_dir, "[0-9][0-9]_seed_*.sql")):
            os.remove(old_f)

    product_tables = []
    extra_ailments = []

    # Kanoniczny porządek dla istniejących dolegliwości, a nowe po nich
    CANONICAL_ORDER = {"sibo": 3, "hashimoto": 4, "insulinoopornosc": 5}
    def sort_key(item):
        code, info = item
        clean = info['table_name'].replace('_produkty', '').lower()
        return (CANONICAL_ORDER.get(clean, 99), clean)

    seed_idx = 3
    for code, info in sorted(discovered.items(), key=sort_key):
        table_name = info['table_name']
        title = info['title']
        file_path = info['file_path']
        clean_name = table_name.replace('_produkty', '')
        # Użyj kanonicznego numeru jeśli znany, w przeciwnym razie inkrementuj
        curr_idx = CANONICAL_ORDER.get(clean_name, seed_idx)
        seed_filename = f"{curr_idx:02d}_seed_{clean_name}.sql"
        output_path = os.path.join(init_dir, seed_filename)

        generate_table_sql(file_path, table_name, title, output_path)
        product_tables.append((table_name, title))
        extra_ailments.append(code)
        seed_idx = max(seed_idx + 1, curr_idx + 1)

    # Dynamiczne generowanie schematu i listy TDP
    generate_schema_sql(os.path.join(init_dir, "01_init_schema.sql"), product_tables)
    generate_tdp_sql(os.path.join(init_dir, "02_seed_tdp.sql"), extra_ailments)

    # Kopiowanie do backend/app/seeds
    if os.path.exists(backend_seeds_dir):
        for f in glob.glob(os.path.join(init_dir, "*.sql")):
            shutil.copy2(f, backend_seeds_dir)
        print(f"Copied all seed scripts to {backend_seeds_dir}")
